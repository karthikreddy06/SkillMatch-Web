"""
SkillMatch Resume Document Parser Service.
Extracts clean, structured plain text from PDF and DOCX documents with robust format handling,
size validation (5MB max), and secure fallback mechanisms.
"""

import io
import re
import base64
import logging
from typing import Tuple, Optional
import requests
from django.conf import settings

import pypdf
import docx

logger = logging.getLogger(__name__)

MAX_RESUME_SIZE = 5 * 1024 * 1024  # 5MB in bytes


class ResumeParserError(Exception):
    """Custom exception raised when resume text extraction fails."""
    def __init__(self, message: str, code: str = 'PARSE_ERROR'):
        super().__init__(message)
        self.message = message
        self.code = code


def _clean_extracted_text(raw_text: str) -> str:
    """Normalize whitespace and line breaks while preserving paragraph boundaries."""
    if not raw_text:
        return ""
    # Replace non-breaking spaces and null bytes
    text = raw_text.replace('\u00a0', ' ').replace('\x00', '')
    # Normalize excessive carriage returns
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    # Collapse 3+ newlines into 2
    text = re.sub(r'\n{3,}', '\n\n', text)
    # Collapse multiple inline spaces/tabs
    text = re.sub(r'[ \t]+', ' ', text)
    return text.strip()


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF byte stream using pypdf."""
    try:
        pdf_stream = io.BytesIO(file_bytes)
        reader = pypdf.PdfReader(pdf_stream)

        if reader.is_encrypted:
            try:
                # Try empty password for unencrypted permissions
                reader.decrypt('')
            except Exception:
                raise ResumeParserError(
                    'The PDF document is password-protected or encrypted. Please remove the password and try again.',
                    code='PASSWORD_PROTECTED'
                )

        pages_text = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                pages_text.append(page_text)

        full_text = '\n\n'.join(pages_text)
        cleaned = _clean_extracted_text(full_text)

        if not cleaned:
            raise ResumeParserError(
                'The PDF document contains no readable text layer (e.g. scanned image or rasterized pages). '
                'Please upload a text-searchable PDF or DOCX.',
                code='EMPTY_DOCUMENT'
            )
        return cleaned

    except ResumeParserError:
        raise
    except Exception as e:
        logger.warning("PDF extraction failed: %s", type(e).__name__)
        raise ResumeParserError(
            'Unable to read PDF document. The file may be damaged or formatted incorrectly.',
            code='CORRUPT_DOCUMENT'
        )


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX byte stream including paragraphs and tables using python-docx."""
    try:
        docx_stream = io.BytesIO(file_bytes)
        doc = docx.Document(docx_stream)

        text_parts = []
        # Paragraphs
        for p in doc.paragraphs:
            p_text = p.text.strip()
            if p_text:
                if p.style and p.style.name and ('bullet' in p.style.name.lower() or 'list' in p.style.name.lower()):
                    if not p_text.startswith(('-', '*', '•', '–')):
                        p_text = f"- {p_text}"
                text_parts.append(p_text)

        # Tables (common in resume formatting)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    text_parts.append(' | '.join(row_cells))

        full_text = '\n'.join(text_parts)
        cleaned = _clean_extracted_text(full_text)

        if not cleaned:
            raise ResumeParserError(
                'The DOCX document contains no readable text. Please check the file contents.',
                code='EMPTY_DOCUMENT'
            )
        return cleaned

    except ResumeParserError:
        raise
    except Exception as e:
        logger.warning("DOCX extraction failed: %s", type(e).__name__)
        raise ResumeParserError(
            'Unable to read Word (.docx) document. The file may be damaged or invalid.',
            code='CORRUPT_DOCUMENT'
        )


def parse_resume_bytes(file_bytes: bytes, filename: str = '') -> Tuple[str, str]:
    """
    Detect document format from bytes or filename, enforce size limit, and extract text.
    Returns tuple of (extracted_text, detected_format).
    """
    if not file_bytes:
        raise ResumeParserError('Empty file received.', code='EMPTY_FILE')

    if len(file_bytes) > MAX_RESUME_SIZE:
        raise ResumeParserError(
            f'File size exceeds maximum allowed limit of 5MB ({len(file_bytes) / (1024 * 1024):.1f}MB).',
            code='FILE_TOO_LARGE'
        )

    # Detect format via magic numbers or filename
    is_pdf = file_bytes.startswith(b'%PDF-')
    is_docx = file_bytes.startswith(b'PK\x03\x04')  # Zip archive header used by docx
    is_doc = file_bytes.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1')  # OLE Compound Document

    ext = (filename.lower().split('.')[-1] if '.' in filename else '')

    if is_pdf or ext == 'pdf':
        return extract_text_from_pdf(file_bytes), 'pdf'
    elif is_docx or ext == 'docx':
        return extract_text_from_docx(file_bytes), 'docx'
    elif is_doc or ext == 'doc':
        raise ResumeParserError(
            'Legacy Word (.doc) format is not supported for automated ATS parsing. '
            'Please convert your resume to PDF or modern Word (.docx) format.',
            code='UNSUPPORTED_FORMAT'
        )
    else:
        # Attempt PDF then DOCX as fallbacks
        try:
            return extract_text_from_pdf(file_bytes), 'pdf'
        except ResumeParserError:
            try:
                return extract_text_from_docx(file_bytes), 'docx'
            except ResumeParserError:
                raise ResumeParserError(
                    'Unsupported document format. Allowed formats: PDF and modern Word (.docx).',
                    code='UNSUPPORTED_FORMAT'
                )


def parse_resume_source(
    file_obj=None,
    resume_url: Optional[str] = None
) -> Tuple[str, str, str]:
    """
    Unified entry point for resume text extraction.
    Supports:
    - Direct UploadedFile from request.FILES
    - Supabase Storage URL
    - Base64 Data URI string

    Returns:
        Tuple of (extracted_text, detected_format, source_name)
    """
    if file_obj:
        filename = getattr(file_obj, 'name', 'uploaded_resume')
        file_bytes = file_obj.read()
        extracted_text, detected_format = parse_resume_bytes(file_bytes, filename)
        return extracted_text, detected_format, filename

    if not resume_url:
        raise ResumeParserError(
            'No resume document provided. Please upload a file or ensure your profile has an attached resume.',
            code='NO_DOCUMENT_SOURCE'
        )

    # Case 1: Base64 Data URI
    if resume_url.startswith('data:'):
        try:
            header, encoded = resume_url.split(',', 1)
            file_bytes = base64.b64decode(encoded)
            filename = 'stored_resume.pdf' if 'pdf' in header else 'stored_resume.docx'
            extracted_text, detected_format = parse_resume_bytes(file_bytes, filename)
            return extracted_text, detected_format, filename
        except ResumeParserError:
            raise
        except Exception:
            raise ResumeParserError(
                'Failed to decode stored base64 resume data.',
                code='CORRUPT_DOCUMENT'
            )

    # Case 2: Remote URL (Supabase storage or external)
    if resume_url.startswith(('http://', 'https://')):
        try:
            # Prepare Supabase service auth if downloading from private bucket
            headers = {}
            service_role_key = getattr(settings, 'SUPABASE_SERVICE_ROLE_KEY', '') or getattr(settings, 'SUPABASE_ANON_KEY', '')
            if service_role_key and 'supabase.co' in resume_url:
                headers['apikey'] = service_role_key
                headers['Authorization'] = f"Bearer {service_role_key}"

            response = requests.get(resume_url, headers=headers, timeout=(5, 20), stream=True)
            if response.status_code == 404:
                raise ResumeParserError(
                    'Stored resume document could not be found in cloud storage (404). Please re-upload your resume.',
                    code='DOCUMENT_NOT_FOUND'
                )
            if response.status_code != 200:
                raise ResumeParserError(
                    f'Unable to download stored resume (HTTP {response.status_code}). Please re-upload your resume.',
                    code='STORAGE_FETCH_FAILED'
                )

            file_bytes = response.content
            # Extract filename from URL
            filename = resume_url.split('/')[-1].split('?')[0] or 'stored_resume'
            extracted_text, detected_format = parse_resume_bytes(file_bytes, filename)
            return extracted_text, detected_format, filename

        except requests.exceptions.Timeout:
            raise ResumeParserError(
                'Timeout while retrieving stored resume from cloud storage. Please try again.',
                code='STORAGE_TIMEOUT'
            )
        except requests.exceptions.RequestException:
            raise ResumeParserError(
                'Network connection error while retrieving stored resume. Please try again.',
                code='NETWORK_ERROR'
            )

    raise ResumeParserError(
        'Invalid or unrecognized resume storage URL format.',
        code='INVALID_SOURCE'
    )
